"""EHS Scorecard — PPTX deck export.

A NEW code path, unlike the PDF: nothing on this platform generated PowerPoint
before. `python-pptx` was already vendored in the environment but was NOT
declared in requirements.txt or pyproject.toml, so an airgapped install would
have shipped without it — both are now declared. It is pure Python (lxml +
Pillow), so it needs no LibreOffice and no headless browser, which is the same
constraint that put the PDF path on fpdf2.

**Same input as the PDF and the dashboard**: `payload.build_payload()`. There is
no export-only query, so a deck cannot disagree with the screen it was exported
from.

⚠ **Format caveat, stated rather than buried.** The build brief referenced "the
Q1 EHS Analysis deck this module was scoped from" as the format to match. That
file is not in this repository and was not supplied. The layout below is a
conventional EHS scorecard deck — title, executive summary, leading, lagging,
trend, per-site, data completeness — built to the design system. It is NOT a
reproduction of the client's deck, because that deck has never been seen. Aligning
slide-for-slide needs the source file.
"""

from __future__ import annotations

import io
from collections.abc import Mapping
from typing import Any

from app.services.display_labels import label

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Midnight Executive. Same values as src/lib/design/midnight.ts and
# app/services/report_pdf.py — one design system, three runtimes.
NAVY = RGBColor(0x0B, 0x1F, 0x4D)
NAVY_MID = RGBColor(0x28, 0x4A, 0x94)
GOLD = RGBColor(0xC9, 0xA9, 0x61)
GOLD_INK = RGBColor(0x7C, 0x5E, 0x07)
ICE = RGBColor(0xE8, 0xEE, 0xF7)
INK = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5B, 0x69, 0x80)
CRIMSON = RGBColor(0x9B, 0x2C, 0x2C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 16:9, the format every projector and every client laptop defaults to.
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _fmt(value: Any, unit: str) -> str:
    """A number, or an em dash — never a zero standing in for absent data."""
    if value is None:
        return "—"
    if unit == "pct":
        return f"{value:g}%"
    if unit in ("rate", "score"):
        return f"{value:g}"
    return f"{int(value):,}"


def _movement(cur: Any, prev: Any, unit: str, good: str) -> tuple[str, RGBColor]:
    if cur is None or prev is None:
        return ("no prior period", MUTED)
    diff = round(float(cur) - float(prev), 2)
    if diff == 0:
        return ("no change", MUTED)
    suffix = "%" if unit == "pct" else ""
    txt = f"{'+' if diff > 0 else '−'}{abs(diff):g}{suffix}"
    if good == "neutral":
        return (txt, MUTED)
    improving = (diff > 0) if good == "up" else (diff < 0)
    return (txt, NAVY_MID if improving else CRIMSON)


def _blank(prs: Presentation):
    # Layout 6 is the blank layout in the default template; everything here is
    # positioned explicitly, so no placeholder inheritance can shift the design.
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, fill: RGBColor):
    from pptx.enum.shapes import MSO_SHAPE

    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _text(slide, x, y, w, h, text, *, size=14, bold=False, color=INK,
          align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Georgia for display type, Calibri for body — the Midnight Executive pairing.
    run.font.name = font
    return box


def _slide_header(slide, title: str, subtitle: str = ""):
    _rect(slide, 0, 0, SLIDE_W, Inches(1.0), NAVY)
    _text(slide, Inches(0.6), Inches(0.16), Inches(9.5), Inches(0.5), title,
          size=24, bold=True, color=WHITE, font="Georgia")
    if subtitle:
        _text(slide, Inches(0.6), Inches(0.62), Inches(9.5), Inches(0.3), subtitle,
              size=11, color=GOLD)


def _table(slide, x, y, w, headers, rows, *, col_widths=None, row_h=Inches(0.32)):
    """A styled table. python-pptx tables inherit a blue Office theme by
    default, so every cell is repainted — an un-restyled table is the fastest
    way for an off-brand colour to reach a client deck."""
    n_rows, n_cols = len(rows) + 1, len(headers)
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, row_h * n_rows)
    tbl = shape.table
    tbl.first_row = True
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(cw))
    for r in range(n_rows):
        tbl.rows[r].height = row_h

    for c, head in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = str(head)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT if c else PP_ALIGN.LEFT
        for run in para.runs:
            run.font.size, run.font.bold = Pt(11), True
            run.font.color.rgb, run.font.name = WHITE, "Calibri"

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            text, colour = (val if isinstance(val, tuple) else (val, INK))
            cell = tbl.cell(r, c)
            cell.text = str(text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else ICE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.RIGHT if c else PP_ALIGN.LEFT
            for run in para.runs:
                run.font.size = Pt(10.5)
                run.font.color.rgb = colour
                run.font.name = "Calibri"
                run.font.bold = c == 0
    return tbl


def render_scorecard_pptx(payload: dict[str, Any], labels: Mapping[str, str] | None = None) -> bytes:
    # `labels`: display-label overrides; each lookup falls through to the literal.
    grain_label = "Monthly" if payload.get("grain") == "month" else "Quarterly"
    site_label = payload.get("siteName") or label(labels, "term.all_sites", "All sites")
    periods = payload.get("periods") or []
    current_label = periods[-1] if periods else "—"

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    # ── 1. Title ────────────────────────────────────────────────────────────
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _rect(s, Inches(0.9), Inches(3.05), Inches(1.6), Inches(0.06), GOLD)
    _text(s, Inches(0.9), Inches(2.1), Inches(11), Inches(0.9), "EHS Scorecard",
          size=44, bold=True, color=WHITE, font="Georgia")
    _text(s, Inches(0.9), Inches(3.3), Inches(11), Inches(0.5),
          f"{grain_label} leading & lagging indicators", size=18, color=GOLD)
    _text(s, Inches(0.9), Inches(4.0), Inches(11), Inches(0.4),
          f"{site_label}   ·   {current_label}", size=15, color=ICE)
    _text(s, Inches(0.9), Inches(6.5), Inches(11), Inches(0.35),
          "SafeOps360  ·  computed from frozen monthly rollups  ·  CONFIDENTIAL",
          size=10, color=RGBColor(0xA7, 0xC3, 0xF8))

    if payload.get("empty"):
        s2 = _blank(prs)
        _slide_header(s2, "No data", current_label)
        _text(s2, Inches(0.6), Inches(1.6), Inches(12), Inches(1),
              payload.get("message") or "No scorecard data for this selection.",
              size=16, color=MUTED)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    cur = payload.get("current") or {}
    prior = payload.get("prior") or {}
    indicators = payload.get("indicators") or []
    by_period = {x["period"]: x for x in payload.get("series", [])}

    # ── 2. Executive summary ────────────────────────────────────────────────
    s = _blank(prs)
    _slide_header(s, "Executive summary", f"{site_label}  ·  {current_label}")
    headline = [i for i in indicators
                if i["key"] in ("observationsLogged", "nearMissReported", "incidentsTotal", "ltifr")]
    card_w, gap = Inches(2.9), Inches(0.28)
    for idx, ind in enumerate(headline):
        x = Inches(0.6) + idx * (card_w + gap)
        _rect(s, x, Inches(1.5), card_w, Inches(1.85), ICE)
        _text(s, x + Inches(0.2), Inches(1.68), card_w - Inches(0.4), Inches(0.3),
              ind["label"].upper(), size=9, bold=True, color=MUTED)
        _text(s, x + Inches(0.2), Inches(2.03), card_w - Inches(0.4), Inches(0.6),
              _fmt(cur.get(ind["key"]), ind["unit"]), size=30, bold=True, color=NAVY,
              font="Georgia")
        txt, colour = _movement(cur.get(ind["key"]), prior.get(ind["key"]),
                                ind["unit"], ind["goodDirection"])
        _text(s, x + Inches(0.2), Inches(2.72), card_w - Inches(0.4), Inches(0.3),
              f"{txt} vs prior period", size=10, color=colour)
    cov = payload.get("sourceCoverage") or {}
    _text(s, Inches(0.6), Inches(3.7), Inches(12.1), Inches(1.2),
          f"Computed from {cov.get('plantMonths', 0)} frozen {label(labels, 'scorecard.plant_months', 'plant-months')} "
          f"({cov.get('firstPeriod') or '—'} to {cov.get('lastPeriod') or '—'}). "
          f"Frequency rates are read from the Manhours module's own submitted figures, not "
          f"recomputed here, so this deck and the manhours return cannot disagree.",
          size=11, color=MUTED)

    # ── 3/4. Leading and lagging tables ─────────────────────────────────────
    for band, title, note in (
        ("leading", "Leading indicators", "What the site is doing about safety"),
        ("lagging", "Lagging indicators", "What happened anyway"),
    ):
        rows_def = [i for i in indicators if i["band"] == band]
        if not rows_def:
            continue
        s = _blank(prs)
        _slide_header(s, title, f"{note}  ·  {site_label}  ·  {current_label}")
        body = []
        for i in rows_def:
            txt, colour = _movement(cur.get(i["key"]), prior.get(i["key"]),
                                    i["unit"], i["goodDirection"])
            body.append([
                i["label"],
                _fmt(cur.get(i["key"]), i["unit"]),
                _fmt(prior.get(i["key"]), i["unit"]),
                (txt, colour),
            ])
        _table(s, Inches(0.6), Inches(1.35), Inches(12.1),
               ["Indicator", current_label, "Prior", "Movement"], body,
               col_widths=[Inches(5.8), Inches(2.1), Inches(2.1), Inches(2.1)])

    # ── 5. Trend ────────────────────────────────────────────────────────────
    s = _blank(prs)
    _slide_header(s, "Trend", f"Last {len(periods)} periods  ·  {site_label}")
    trend_keys = ["observationsLogged", "nearMissReported", "incidentsTotal", "ltiCount", "ltifr"]
    trend_def = [i for i in indicators if i["key"] in trend_keys]
    shown = periods[-8:]
    body = [[i["label"]] + [_fmt(by_period.get(p, {}).get(i["key"]), i["unit"]) for p in shown]
            for i in trend_def]
    first_w = Inches(3.6)
    rest = (Inches(12.1) - first_w) / max(len(shown), 1)
    _table(s, Inches(0.6), Inches(1.35), Inches(12.1), ["Indicator"] + shown, body,
           col_widths=[first_w] + [rest] * len(shown))
    if len(periods) > len(shown):
        _text(s, Inches(0.6), Inches(6.6), Inches(12), Inches(0.3),
              f"Showing the most recent {len(shown)} of {len(periods)} periods.",
              size=10, color=MUTED)

    # ── 6. By site ──────────────────────────────────────────────────────────
    by_site = payload.get("bySite") or []
    if len(by_site) > 1:
        s = _blank(prs)
        _slide_header(s, label(labels, "scorecard.by_site", "By site"), current_label)
        body = [[
            (x.get("siteName") or "")[:44],
            _fmt(x.get("observationsLogged"), "count"),
            _fmt(x.get("nearMissReported"), "count"),
            _fmt(x.get("incidentsTotal"), "count"),
            _fmt(x.get("ltiCount"), "count"),
            _fmt(x.get("ltifr"), "rate"),
        ] for x in by_site[:14]]
        _table(s, Inches(0.6), Inches(1.35), Inches(12.1),
               [label(labels, "term.site", "Site"), "Observations", "Near miss", "Incidents", "LTI", "LTIFR"], body,
               col_widths=[Inches(4.9), Inches(1.6), Inches(1.5), Inches(1.5),
                           Inches(1.3), Inches(1.3)])

    # ── 7. Data completeness. Last slide, never omitted. ─────────────────────
    s = _blank(prs)
    _slide_header(s, "Data completeness", "What this scorecard could not measure")
    gaps = payload.get("gaps") or []
    if not gaps:
        _text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.5),
              "Every indicator on this scorecard was computable for the selected period.",
              size=14, color=INK)
    else:
        _text(s, Inches(0.6), Inches(1.35), Inches(12.1), Inches(0.6),
              "The indicators below could not be computed. A blank or a zero against them is "
              "an absence of data, not a result.", size=12, color=MUTED)
        y = Inches(2.0)
        for g in gaps[:8]:
            _rect(s, Inches(0.6), y, Inches(0.05), Inches(0.52), GOLD)
            _text(s, Inches(0.85), y, Inches(11.8), Inches(0.26),
                  g.get("indicator", "Indicator"), size=12, bold=True, color=GOLD_INK)
            _text(s, Inches(0.85), y + Inches(0.26), Inches(11.8), Inches(0.26),
                  g.get("reason", ""), size=10, color=MUTED)
            y += Inches(0.62)
    if cov.get("lastPeriodWithExposure") and cov.get("lastPeriod") != cov.get("lastPeriodWithExposure"):
        _text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.6),
              f"Exposure data (manhours) is recorded only to {cov['lastPeriodWithExposure']}, "
              f"while activity data runs to {cov['lastPeriod']} — frequency rates are "
              f"unavailable for the most recent periods, not zero.",
              size=10, color=GOLD_INK)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


__all__ = ["render_scorecard_pptx"]
