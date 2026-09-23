"""Export regression baseline — fetch every read-only PDF/PPTX export per persona
from a running API, store the files, and extract normalised text for diffing.

    python scripts/verify/export_baseline.py capture <outDir> [--personas mfg,apparel]
    python scripts/verify/export_baseline.py diff <beforeDir> <afterDir>

Only GET endpoints that do not write are called (the audit-compliance report PDF
stamps pdfAttachmentId, so it is excluded). Dates, times and digits are masked
in the extracted text so a re-render on another day is not reported as drift.
A response that is not the requested format (e.g. CSV served for a PDF request)
is recorded as a FORMAT MISMATCH and fails the diff.
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
from pathlib import Path

import httpx
import psycopg2
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[2]
load_dotenv(ROOT / ".env")
API = os.environ.get("API_BASE", "http://localhost:8000")
PW = os.environ.get("DEMO_PW", "demo123")

PERSONAS = {
    "mfg": ("priya.nair@safeops360.in", "NW"),
    "mfg-cro": ("anand.krishnan@safeops360.in", "NW"),
    "apparel": ("harpreet.singh@meridian-apparel.in", "MAG-LDH"),
    "industry": ("plant-head.acs@safeops360.in", "ACS"),
    "retail": ("store-ops.admin@meridian-retail.in", "MR-DC01"),
}
ERM_KINDS = ["register", "assessments", "treatments", "overdue", "department", "heatmap", "escalations", "acceptances"]


def _db():
    url = os.environ["DATABASE_URL_SYNC"]
    url = re.sub(r"^postgresql\+\w+://", "postgresql://", url)
    return psycopg2.connect(url)


def _first(cur, sql: str, *args):
    cur.execute(sql, args)
    row = cur.fetchone()
    return row[0] if row else None


def endpoints(plant_code: str) -> list[tuple[str, str, str]]:
    """(name, path, expected kind) for this persona's home plant."""
    with _db() as c, c.cursor() as cur:
        plant_id = _first(cur, 'select id from "Plant" where code=%s', plant_code)
        permit = _first(cur, 'select id from "Permit" where "plantId"=%s order by "createdAt" limit 1', plant_id)
        form = None
        try:
            cur.execute(
                'select f."incidentId", f.id from "StatutoryFormInstance" f join "Incident" i on i.id=f."incidentId" '
                'where i."plantId"=%s order by f."createdAt" limit 1', (plant_id,))
            form = cur.fetchone()
        except Exception:
            c.rollback()
        brsr = None
        try:
            brsr = _first(cur, 'select id from "BrsrCycle" order by "createdAt" limit 1')
        except Exception:
            c.rollback()
    eps = [
        ("scorecard.pdf", f"/api/scorecard/export.pdf?site={plant_id}", "pdf"),
        ("scorecard.pptx", f"/api/scorecard/export.pptx?site={plant_id}", "pptx"),
        ("kaizen-register.pdf", f"/api/be/kaizen/export?format=pdf&plantId={plant_id}", "pdf"),
    ]
    eps += [(f"erm-{k}.pdf", f"/api/erm/reports/{k}.pdf", "pdf") for k in ERM_KINDS]
    if permit:
        eps.append(("ptw-closeout.pdf", f"/api/ptw/{permit}/report", "pdf"))
    if form:
        eps.append(("statutory-form.pdf", f"/api/incidents/{form[0]}/statutory-forms/{form[1]}/download", "pdf"))
    if brsr:
        eps.append(("brsr-report.pdf", f"/api/brsr/cycles/{brsr}/report.pdf", "pdf"))
    return eps


def _mask(text: str) -> str:
    text = re.sub(r"\d{1,4}[-/ ](\d{1,2}|[A-Za-z]{3,9})[-/ ,]+\d{2,4}", "<date>", text)
    text = re.sub(r"\d{1,2}:\d{2}(:\d{2})?", "<time>", text)
    text = re.sub(r"\d+", "#", text)
    return "\n".join(line.rstrip() for line in text.splitlines() if line.strip())


def extract(kind: str, data: bytes) -> str:
    if kind == "pdf":
        if not data.startswith(b"%PDF"):
            return f"FORMAT MISMATCH: expected PDF, got {data[:40]!r}"
        from pypdf import PdfReader
        return _mask("\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages))
    if kind == "pptx":
        if not data.startswith(b"PK"):
            return f"FORMAT MISMATCH: expected PPTX, got {data[:40]!r}"
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            out.append(f"--- slide {i}")
            for sh in slide.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        out.append(" | ".join(c.text for c in row.cells))
        return _mask("\n".join(out))
    return ""


def capture(out_dir: Path, keys: list[str]) -> None:
    for key in keys:
        email, plant = PERSONAS[key]
        d = out_dir / key
        d.mkdir(parents=True, exist_ok=True)
        with httpx.Client(base_url=API, timeout=180) as http:
            r = http.post("/api/auth/login", json={"email": email, "password": PW})
            r.raise_for_status()
            token = r.json().get("accessToken") or r.json().get("access_token") or r.json().get("token")
            http.headers["Authorization"] = f"Bearer {token}"
            summary = {}
            for name, path, kind in endpoints(plant):
                resp = http.get(path)
                summary[name] = {"path": path, "status": resp.status_code,
                                 "contentType": resp.headers.get("content-type"), "bytes": len(resp.content)}
                if resp.status_code == 200:
                    (d / name).write_bytes(resp.content)
                    (d / f"{name}.txt").write_text(extract(kind, resp.content), encoding="utf-8")
                print(f"[{key}] {name}: {resp.status_code} {summary[name]['contentType']} {len(resp.content)}B")
            (d / "summary.json").write_text(json.dumps(summary, indent=1), encoding="utf-8")


def diff(before: Path, after: Path) -> int:
    import difflib
    drift = 0
    for sfile in sorted(before.glob("*/summary.json")):
        key = sfile.parent.name
        b = json.loads(sfile.read_text())
        af = after / key / "summary.json"
        a = json.loads(af.read_text()) if af.exists() else {}
        for name, meta in b.items():
            am = a.get(name)
            if not am or am["status"] != meta["status"] or am["contentType"] != meta["contentType"]:
                drift += 1
                print(f"[{key}] {name}: status/type {meta['status']} {meta['contentType']} -> "
                      f"{am and am['status']} {am and am['contentType']}")
                continue
            bt, at = before / key / f"{name}.txt", after / key / f"{name}.txt"
            if bt.exists() and at.exists():
                bl, al = bt.read_text(encoding="utf-8").splitlines(), at.read_text(encoding="utf-8").splitlines()
                if "FORMAT MISMATCH" in at.read_text(encoding="utf-8")[:40]:
                    drift += 1
                    print(f"[{key}] {name}: {al[0]}")
                elif bl != al:
                    drift += 1
                    print(f"[{key}] {name}: text differs")
                    for line in list(difflib.unified_diff(bl, al, lineterm="", n=0))[2:40]:
                        print("    " + line)
    print(f"\n{drift} differing exports" if drift else "\nNo export drift.")
    return 1 if drift else 0


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "capture":
        keys = (sys.argv[4] if len(sys.argv) > 4 and sys.argv[3] == "--personas" else "mfg,mfg-cro,apparel,industry").split(",")
        capture(Path(sys.argv[2]), keys)
    else:
        sys.exit(diff(Path(sys.argv[2]), Path(sys.argv[3])))
