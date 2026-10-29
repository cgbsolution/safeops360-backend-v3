"""Step 1 (HIRA): read a QSD/24/F HIRA workbook with its QSD/25/F control sheet.

    python -I parse_hira.py <HIRA.xlsx> <hira_raw.json>

Sheet "HIRA " (note the trailing space in the plant's file): from row 17,
columns A-T = S.No, Activity, R/NR, Hazard, OH&S risk, Control reference,
LC, DC, CE, EC, SP, LH, CN, SV, DE, GT, significant score, previous GT,
acceptable, legal reference. Section titles are merged rows.
Sheet "Controls": from row 10, columns A-M = S.No, Activity, Hazard, Risk,
Elimination, Tech modification, Procedure, Engineering, Signage,
Administrative, Competence, PIP, PPE.
Merged cells are filled down; team names in the header rows are not read.
"""

from __future__ import annotations

import json
import sys

import openpyxl


def grid(ws):
    g = {}
    for row in ws.iter_rows():
        for c in row:
            g[(c.row, c.column)] = c.value
    for mr in ws.merged_cells.ranges:
        v = g.get((mr.min_row, mr.min_col))
        for r in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                g[(r, cc)] = v
    return g, ws.max_row


def s(v):
    if isinstance(v, str):
        v = " ".join(v.split())
        return v or None
    return v


FOOTER = ("1)", "Prepared", "Approved", "Name", "Signature", "Reviewed", "*")


def main() -> int:
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True)
    hira_ws = next(ws for ws in wb.worksheets if ws.title.strip().upper() == "HIRA")
    h, mx = grid(hira_ws)
    hira, section, last = [], None, None
    for r in range(17, mx + 1):
        row = [s(h.get((r, c))) for c in range(1, 24)]
        a = row[0]
        if isinstance(a, str) and all(x in (None, a) for x in row[1:6]):
            if not a.startswith(FOOTER):
                section, last = a, None
            continue
        if row[3] is None and row[1] is None:
            continue
        act = row[1] or last
        last = act
        hira.append(dict(xl_row=r, section=section, sno=row[0], activity=act, rnr=row[2], hazard=row[3], risk=row[4],
                         control_ref=row[5], LC=row[6], DC=row[7], CE=row[8], EC=row[9], SP=row[10], LH=row[11], CN=row[12],
                         SV=row[13], DE=row[14], GT=row[15], sig=row[16], prev=row[17], acceptable=row[18], legal=row[19]))
    c, mx2 = grid(wb["Controls"])
    cols = ["elimination", "tech_modification", "procedure", "engineering", "signage", "administrative", "competence", "pip", "ppe"]
    ctrl, section, last = [], None, None
    for r in range(10, mx2 + 1):
        row = [s(c.get((r, k))) for k in range(1, 14)]
        a = row[0]
        if isinstance(a, str) and all(x in (None, a) for x in row[1:5]):
            if not a.startswith(FOOTER):
                section, last = a, None
            continue
        if row[2] is None and row[1] is None:
            continue
        act = row[1] or last
        last = act
        d = dict(xl_row=r, section=section, sno=row[0], activity=act, hazard=row[2], risk=row[3])
        d.update(zip(cols, row[4:13]))
        ctrl.append(d)
    json.dump({"hira": hira, "controls": ctrl}, open(out, "w", encoding="utf-8"), indent=1, default=str)
    print(f"{len(hira)} HIRA rows, {len(ctrl)} control rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
