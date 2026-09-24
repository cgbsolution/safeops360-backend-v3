"""Step 4 — diagnostic verification of the Meridian Retail tenant (live API + DB).

Runs against the running backend (API_BASE, default http://localhost:8000) as
real users; nothing is written. Every check prints PASS/FAIL with evidence.

    python -m scripts.meridian_retail.verify
"""

from __future__ import annotations

import io
import os
import re
import sys

import httpx

from scripts.meridian_retail.common import EMAIL_DOMAIN, PASSWORD, conn

API = os.environ.get("API_BASE", "http://localhost:8000")
RETAIL_ADMIN = f"store-ops.admin@{EMAIL_DOMAIN}"
MFG_USER = "priya.nair@safeops360.in"
LEAK = re.compile(r"\b(Plants?|Factory|Factories|Shift Supervisor|Plant Head)\b")
results: list[tuple[bool, str]] = []


def check(ok: bool, what: str, evidence: str = "") -> None:
    results.append((ok, what))
    print(f"{'PASS' if ok else 'FAIL'}  {what}" + (f"  — {evidence}" if evidence else ""))


def login(email: str) -> httpx.Client:
    r = httpx.post(f"{API}/api/auth/login", json={"email": email, "password": PASSWORD}, timeout=60)
    r.raise_for_status()
    body = r.json()
    c = httpx.Client(base_url=API, timeout=180, follow_redirects=True,
                     headers={"Authorization": f"Bearer {body['access_token']}"})
    return c


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)


def pptx_text(data: bytes) -> str:
    from pptx import Presentation
    out = []
    for s in Presentation(io.BytesIO(data)).slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
    return "\n".join(out)


def main() -> int:
    db = conn()
    cur = db.cursor()
    cur.execute('select id, code from "Plant" where code like %s', ("MR-%",))
    retail = dict(cur.fetchall())
    retail_ids = set(retail)
    cur.execute('select id from "Plant" where code in (%s,%s)', ("NW", "SW"))
    mfg_ids = {r[0] for r in cur.fetchall()}
    check(len(retail) == 43, "43 Retail sites (40 stores + 3 DCs)", f"{len(retail)} found")

    # ── Isolation: zero shared rows ──
    cur.execute('select count(*) from "User" where email like %s and "plantId" <> all(%s)', (f"%@{EMAIL_DOMAIN}", list(retail_ids)))
    check(cur.fetchone()[0] == 0, "every Retail user's home plant is a Retail site")
    cur.execute('select count(*) from "UserRole" ur join "User" u on u.id=ur."userId" where u.email like %s '
                'and ur."scopeType"=%s and ur."scopeValue" <> all(%s)', (f"%@{EMAIL_DOMAIN}", "PLANT", list(retail_ids)))
    check(cur.fetchone()[0] == 0, "no Retail user holds a role scoped to a non-Retail site")
    cur.execute('select count(*) from "RolePermission" rp join "Role" r on r.id=rp."roleId" where r.code like %s and rp.scope=%s',
                ("RETAIL_%", "ALL_PLANTS"))
    check(cur.fetchone()[0] == 0, "Retail roles hold zero ALL_PLANTS grants")
    cur.execute('select count(*) from "UserRole" ur join "User" u on u.id=ur."userId" join "Role" r on r.id=ur."roleId" '
                'where u.email like %s and r.code not like %s', (f"%@{EMAIL_DOMAIN}", "RETAIL_%"))
    check(cur.fetchone()[0] == 0, "Retail users hold only RETAIL_* roles")
    for table, col in (("Incident", "plantId"), ("NearMiss", "plantId"), ("Permit", "plantId"), ("CaptureSubmission", "plantId"),
                       ("FireEquipment", "plantId"), ("CamsEngagement", "siteId")):
        cur.execute(f'select count(*) from "{table}" where "{col}" = any(%s)', (list(mfg_ids),))
        mfg_n = cur.fetchone()[0]
        cur.execute(f'select count(*) from "{table}" where "{col}" = any(%s)', (list(retail_ids),))
        ret_n = cur.fetchone()[0]
        check(ret_n > 0, f"{table}: Retail rows exist, on Retail sites only", f"retail={ret_n}, manufacturing={mfg_n} (disjoint by site)")

    ra = login(RETAIL_ADMIN)
    mu = login(MFG_USER)
    for path, key in (("/api/incidents", "plantId"), ("/api/near-miss", "plantId"), ("/api/ptw", "plantId"),
                      ("/api/cams/unified-engagements", "siteId"), ("/api/fire/equipment", "plantId")):
        r = ra.get(path)
        items = r.json() if isinstance(r.json(), list) else (r.json().get("items") or r.json().get("data") or [])
        foreign = [i for i in items if isinstance(i, dict) and i.get(key) and i[key] not in retail_ids]
        check(r.status_code == 200 and not foreign, f"Retail admin sees only Retail rows: {path}",
              f"{len(items)} rows, {len(foreign)} foreign")
    r = mu.get("/api/epc/sites/")
    sites = r.json() if isinstance(r.json(), list) else (r.json().get("sites") or r.json().get("items") or [])
    check(not any(str(s.get("siteCode", "")).startswith("MR-") for s in sites), "Manufacturing user sees no Retail EPC projects",
          f"{len(sites)} sites")
    r = ra.get("/api/epc/sites/")
    sites = r.json() if isinstance(r.json(), list) else (r.json().get("sites") or r.json().get("items") or [])
    check(sites and all(str(s.get("siteCode", "")).startswith("MR-") for s in sites), "Retail admin sees only Retail EPC projects",
          f"{len(sites)} sites")

    # ── Labels ──
    store = next(pid for pid, code in retail.items() if code == "MR-S001")
    lbl = ra.get("/api/display-labels", params={"plantId": store}).json()["labels"]
    for k, v in (("term.plant", "Store"), ("term.factory", "Distribution Center"), ("term.shift_supervisor", "Store Manager"),
                 ("nav./ptw", "Facilities & DC Maintenance Permits")):
        check(lbl.get(k) == v, f"label {k} → {v}", f"got {lbl.get(k)!r}")
    nw = next(iter(mfg_ids))
    check(mu.get("/api/display-labels", params={"plantId": nw}).json()["labels"] == {},
          "Manufacturing sites have no overrides (fall back to defaults)")

    # ── PTW permit-type curation (Step 6) ──
    curated = ["HOT_WORK", "WORK_AT_HEIGHT", "ELECTRICAL_LOTO", "LIFTING", "GENERAL_COLD"]
    cfgs = ra.get("/api/ptw-type-config").json()["configs"]
    check(set(cfgs) == retail_ids and all(c["enabledTypes"] == curated for c in cfgs.values()),
          "every Retail site offers exactly the 5 curated permit types (no Confined Space / Excavation)",
          f"{len(cfgs)} sites configured")
    check(all(c["defaultType"] == "ELECTRICAL_LOTO" for c in cfgs.values()),
          "PTW wizard opens on Electrical / LOTO at Retail sites, not Hot Work")
    check(all(set(c["blockedHazards"]) == {"CONFINED_SPACE", "EXCAVATION"} for c in cfgs.values()),
          "Confined Space / Excavation annexures also hidden at Retail sites")
    cur.execute('select count(*) from "PlantPermitTypeConfig" where "plantId" <> all(%s)', (list(retail_ids),))
    check(cur.fetchone()[0] == 0, "no non-Retail site has curation (full type set, Hot Work default)")
    plants_seen = {p["id"] for p in ra.get("/api/plants").json()}
    check(plants_seen and plants_seen <= retail_ids, "PTW wizard plant picker offers Retail sites only",
          f"{len(plants_seen)} plants")
    cur.execute('select id from "User" where email like %s and email <> %s order by email limit 2',
                (f"%@{EMAIL_DOMAIN}", RETAIL_ADMIN))
    iss, rec = (r[0] for r in cur.fetchall())
    for bad in ("CONFINED_SPACE", "EXCAVATION"):
        r = ra.post("/api/ptw", json={"type": bad, "plantId": store, "location": "Stockroom",
                                      "scopeOfWork": "Verification probe - must be refused",
                                      "validFrom": "2026-10-01T08:00:00Z", "validTo": "2026-10-01T12:00:00Z",
                                      "issuerId": iss, "receiverId": rec})
        check(r.status_code == 400 and "not used at this site" in r.text, f"{bad} permit refused server-side at a Retail site",
              f"HTTP {r.status_code}")
    cur.execute('select type, count(*) from "Permit" where "plantId" = any(%s) group by type', (list(retail_ids),))
    ptypes = dict(cur.fetchall())
    check(10 <= sum(ptypes.values()) <= 12 and set(ptypes) <= set(curated), "10–12 Retail PTW records, curated types only",
          str(ptypes))

    # ── Module switches ──
    mods = ra.get("/api/licensing/modules", params={"plantId": store}).json()
    enabled, disabled = set(mods["enabledModules"]), set(mods.get("disabledModules", []))
    off = {"MOC", "ERM", "BCM"} - enabled
    check(off == {"MOC", "ERM", "BCM"}, "MOC / ERM (and BCM) off at Retail sites", f"still on: {({'MOC','ERM','BCM'} & enabled) or '-'}")
    check({"BRSR", "BUSINESS_EXCELLENCE", "TRAINING_ENGINE", "CAMS_GENERAL"} <= disabled,
          "BRSR / Business Excellence / training rule-engine / general CAMS off at Retail sites", f"disabled={sorted(disabled)}")
    check({"CAMS", "PTW", "INCIDENT", "NEAR_MISS", "EPC"} <= enabled and not ({"ALERTS", "FIRE", "LOTO", "CAPTURE", "SCORECARD"} & disabled),
          "Daily Brief, Fire, CAMS, CSM(EPC), Field Reports, PTW/LOTO, Scorecard on at Retail sites")
    hdr = {"X-Active-Plant": store}
    for path in ("/api/brsr/cycles", "/api/moc/change-requests", "/api/erm/risks", "/api/be/kaizen", "/api/audit-compliance/audits"):
        code = ra.get(path, headers=hdr).status_code
        check(code == 403, f"disabled module API refused at a Retail site: {path}", f"HTTP {code}")

    # ── CAMS Fire-only ──
    code = ra.post("/api/cams/engagements", headers=hdr, json={"title": "x", "engagementType": "INTERNAL_AUDIT",
                                                              "siteId": store, "leadAuditorId": "x", "plannedDate": "2026-10-01T00:00:00Z"}).status_code
    check(code == 403, "generic (non-fire) CAMS engagement creation refused at a Retail site", f"HTTP {code}")
    engs = ra.get("/api/cams/unified-engagements").json()["items"]
    check(engs and all(e.get("sourceModule") == "FIRE" for e in engs), "every CAMS engagement a Retail user can see is a Fire one",
          f"{len(engs)} engagements")
    audits = ra.get("/api/fire/audits").json()["items"]
    done = [a for a in audits if a["status"] in ("REPORT_ISSUED", "CLOSED")]
    check(len(done) >= 2, "2–3 completed CAMS Fire Safety audits", f"{len(done)} completed, {len(audits)} total")
    cur.execute('select count(*), count(distinct "subjectUserId") from "IndependenceEvent" where "engagementCode" like %s', ("FSA-MR-%",))
    ev = cur.fetchone()
    check(ev[0] > 0, "independence engine assignment recorded for the Fire Safety audits", f"{ev[0]} events, {ev[1]} auditors")
    if done:
        snap = ra.get(f"/api/fire/compliance/engagement/{done[0]['id']}").json()
        reg, comp = snap.get("register") or {}, snap.get("overall") or {}
        cur.execute('select count(*) from "CamsEngagementAsset" where "engagementId"=%s', (done[0]["id"],))
        linked = cur.fetchone()[0]
        check(reg.get("total") == linked and comp.get("owed", 0) > 0 and {s["via"] for s in snap.get("sources", [])}
              == {"Fire register", "Routine checklists"},
              "Compliance Snapshot computed live from register + checklist completion, with via-badges",
              f"assets={reg.get('total')} (linked {linked}), owed={comp.get('owed')}, completed={comp.get('completed')}, rate={comp.get('rate')}")

    # ── Register: counts, mixed history ──
    cur.execute('select type, count(*) from "FireEquipment" where "plantId" = any(%s) group by type', (list(retail_ids),))
    by_type = dict(cur.fetchall())
    total = sum(by_type.values())
    check(120 <= total <= 150 and set(by_type) >= {"FIRE_EXTINGUISHER", "FIRE_ALARM_PANEL", "FIRE_HYDRANT_SYSTEM"},
          "120–150 fire assets across extinguisher / alarm / hydrant", f"{total}: {by_type}")
    cur.execute('select status, count(*) from "FireEquipment" where "plantId" = any(%s) group by status', (list(retail_ids),))
    st = dict(cur.fetchall())
    check(st.get("OVERDUE", 0) > 0 and st.get("ACTIVE", 0) > 0, "register is not all-green (overdue + on-schedule)", str(st))

    # ── Exports render with override labels (not CSV) ──
    for slug in ("extinguisher-register", "alarm-panel-register", "hydrant-system-register"):
        r = ra.get(f"/api/fire/registers/{slug}/export.pdf")
        ok = r.status_code == 200 and r.headers.get("content-type", "").startswith("application/pdf")
        t = pdf_text(r.content) if ok else ""
        check(ok and "Meridian Retail" in t and "Store" in t, f"fire register PDF renders with Retail branding: {slug}",
              f"HTTP {r.status_code} {r.headers.get('content-type')}, {len(r.content)} bytes")
    for ext, reader in (("pdf", pdf_text), ("pptx", pptx_text)):
        r = ra.get(f"/api/scorecard/export.{ext}")
        ok = r.status_code == 200 and ("pdf" in r.headers.get("content-type", "") or "presentation" in r.headers.get("content-type", ""))
        t = reader(r.content) if ok else ""
        leaks = sorted(set(LEAK.findall(t)))
        check(ok and "store" in t.lower() and not leaks, f"EHS Scorecard {ext.upper()} renders with Store labels, no Plant/Factory",
              f"HTTP {r.status_code}, leaks={leaks}")
        if ext == "pdf":
            check("could NOT be computed" in t or "not a zero" in t.lower() or "no manhours" in t.lower(),
                  "scorecard shows 'cannot be computed' for missing manhours (not zero)")

    failed = [w for ok, w in results if not ok]
    print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
